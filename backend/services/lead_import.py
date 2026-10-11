"""Bounded spreadsheet import: preview, skip duplicates, then insert with audit."""
import base64
import binascii
import csv
import hashlib
import io
import json
import unicodedata
import uuid
import zipfile
from datetime import date, datetime
from itertools import islice
from pathlib import Path

from fastapi import HTTPException
from pydantic import Field, ValidationError
from sqlalchemy import select, text
from starlette.concurrency import run_in_threadpool

from database.lead_models import Lead
from models.lead_schemas import StrictInput, LeadInput, ContactInput
from services.lead_service import record_changes, save_contact

MAX_BYTES = 3 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 60


class SpreadsheetInput(StrictInput):
    filename: str = Field(min_length=1, max_length=255)
    content: str = Field(min_length=1, max_length=4 * ((MAX_BYTES + 2) // 3))
    preview_digest: str | None = Field(default=None, pattern=r'^[a-f0-9]{64}$')


def normalized(value):
    value = ''.join(c for c in unicodedata.normalize('NFKD', str(value)) if not unicodedata.combining(c))
    return ' '.join(value.lower().replace('_', ' ').replace('-', ' ').split())


LABELS = {
    'name': 'Název', 'type': 'Typ', 'type_label': 'Zobrazovaný typ', 'city': 'Město',
    'region': 'Kraj', 'address': 'Adresa', 'country': 'Země', 'postal_code': 'PSČ',
    'district': 'Okres', 'lat': 'Zeměpisná šířka', 'lng': 'Zeměpisná délka',
    'location_accuracy': 'Přesnost polohy', 'cluster': 'Klastr', 'plan': 'Tarif',
    'programs': 'Programy', 'catalog_programs': 'Počet programů v katalogu',
    'booking': 'Způsob rezervace', 'online': 'Online termíny',
    'online_completion': 'Dokončení rezervace online', 'activity': 'Obchodní zařazení',
    'relationship': 'Dosavadní vztah', 'regularity': 'Pravidelnost nabídky',
    'score': 'Potenciál', 'why': 'Důvod oslovení', 'catalog': 'Katalogová hodnota',
    'catalog_role': 'Role pro katalog', 'website_url': 'Web', 'program_url': 'Odkaz na programy',
    'stage': 'Stav', 'result_reason': 'Důvod výsledku', 'next_action': 'Další krok',
    'next_action_date': 'Termín dalšího kroku', 'last_contact': 'Poslední kontakt',
    'internal_note': 'Interní poznámka', 'contact_first_name': 'Jméno kontaktu',
    'contact_last_name': 'Příjmení kontaktu', 'contact_position': 'Pozice kontaktu',
    'contact_email': 'E-mail', 'contact_phone': 'Telefon', 'contact_note': 'Poznámka ke kontaktu',
}
ALIASES = {normalized(alias): field for field, label in LABELS.items() for alias in (field, label)}
ALIASES.update({normalized(alias): field for field, aliases in {
    'name': ['Instituce', 'Název instituce'], 'type': ['Typ instituce'],
    'plan': ['Plán'], 'website_url': ['Webová stránka', 'URL'],
    'contact_email': ['Email', 'E-mail kontaktu'], 'contact_phone': ['Telefon kontaktu'],
    'internal_note': ['Poznámka'],
}.items() for alias in aliases})
TYPES = {'muzeum': 'museum', 'galerie': 'gallery', 'knihovna': 'library',
         'botanicka zahrada': 'botanical', 'science centrum': 'science',
         'navstevnicke centrum': 'visitor', 'divadlo': 'theatre', 'zoo': 'zoo',
         'ekocentrum': 'eco', 'vzdelavaci instituce': 'education', 'kino': 'cinema'}


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(data):
    try:
        content = base64.b64decode(data.content, validate=True)
    except (ValueError, binascii.Error):
        raise HTTPException(422, 'Soubor se nepodařilo načíst. Vyberte jej znovu.')
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Soubor může mít nejvýše 3 MB.')
    extension = Path(data.filename).suffix.lower()
    try:
        if extension == '.csv':
            try:
                decoded = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                decoded = content.decode('cp1250')
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',;\t')
                delimiter = dialect.delimiter
            except csv.Error:
                # A single-column file with institution names is valid too.
                delimiter = ';' if ';' in decoded.splitlines()[0] else ','
            rows = list(islice(csv.reader(io.StringIO(decoded), delimiter=delimiter, strict=True), MAX_ROWS + 2))
        elif extension == '.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000 or sum(x.file_size for x in archive.infolist()) > 20 * 1024 * 1024:
                    raise HTTPException(413, 'Rozbalený Excel je příliš velký. Použijte menší tabulku.')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book.worksheets[0]
                if (sheet.max_column or 0) > MAX_COLUMNS:
                    raise HTTPException(422, 'Tabulka může mít nejvýše 60 sloupců.')
                rows = []
                # Ignore misleading dimensions but always bound actual iteration.
                sheet.reset_dimensions()
                for cells in islice(sheet.iter_rows(max_col=MAX_COLUMNS + 1), MAX_ROWS + 2):
                    if any(c.data_type == 'f' for c in cells):
                        raise HTTPException(422, 'Excel obsahuje vzorce. Vložte jejich hodnoty a nahrajte soubor znovu.')
                    values = [c.value for c in cells]
                    while values and values[-1] is None:
                        values.pop()
                    rows.append(values)
            finally:
                book.close()
        elif extension == '.xls':
            import xlrd
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                sheet = book.sheet_by_index(0)
                if sheet.ncols > MAX_COLUMNS or sheet.nrows > MAX_ROWS + 1:
                    raise HTTPException(422, 'Tabulka může mít nejvýše 500 řádků a 60 sloupců.')
                rows = [[xlrd.xldate.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value
                         for c in sheet.row(i)] for i in range(sheet.nrows)]
            finally:
                book.release_resources()
        else:
            raise HTTPException(422, 'Vyberte soubor CSV, XLSX nebo XLS.')
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Soubor nelze přečíst. Zkontrolujte formát a první list tabulky.')
    if len(rows) > MAX_ROWS + 1 or any(len(row) > MAX_COLUMNS for row in rows):
        raise HTTPException(422, 'Tabulka může mít nejvýše 500 řádků a 60 sloupců.')
    if not rows or not any(cell_text(x) for x in rows[0]):
        raise HTTPException(422, 'Chybí řádek s názvy sloupců.')
    if any(len(cell_text(c)) > 10000 for row in rows for c in row):
        raise HTTPException(422, 'Jedna buňka může mít nejvýše 10 000 znaků.')
    return rows


def parse_sheet(data):
    rows = read_rows(data)
    headers = [cell_text(c) for c in rows[0]]
    fields = [ALIASES.get(normalized(h)) for h in headers]
    if 'name' not in fields:
        raise HTTPException(422, 'Chybí sloupec Název (nebo name / Název instituce).')
    known = [f for f in fields if f]
    if len(known) != len(set(known)):
        raise HTTPException(422, 'Více sloupců odpovídá stejnému poli. Ponechte každý sloupec pouze jednou.')
    result = []
    for number, row in enumerate(rows[1:], 2):
        if not any(cell_text(c) for c in row):
            continue
        values = {field: cell_text(row[i]) for i, field in enumerate(fields) if field and i < len(row) and cell_text(row[i])}
        entry = {'row': number, 'name': values.get('name', ''), 'city': values.get('city', ''), 'status': 'new', 'message': ''}
        try:
            if len(row) > len(headers):
                raise ValueError('Řádek má více buněk než hlavička.')
            contact = {key: values.pop('contact_' + key) for key in ContactInput.model_fields if 'contact_' + key in values}
            original_type = values.get('type', '')
            values['type'] = TYPES.get(normalized(original_type), original_type or 'other')
            if normalized(original_type) in TYPES or not original_type:
                values.setdefault('type_label', original_type or 'Instituce')
            for key in ('lat', 'lng'):
                if key in values:
                    values[key] = float(values[key].replace(',', '.'))
            for key in ('next_action_date', 'last_contact'):
                if key in values and '.' in values[key]:
                    values[key] = datetime.strptime(values[key].replace(' ', ''), '%d.%m.%Y').date()
            # Stage/plan matching tolerates accents, spaces and dash variants.
            from models.lead_schemas import STAGES, PLANS
            for key, options in [('stage', STAGES), ('plan', PLANS)]:
                if key in values:
                    match = next((o for o in options if normalized(o.replace('–', '-')) == normalized(values[key].replace('–', '-'))), None)
                    if match:
                        values[key] = match
            entry['data'] = LeadInput(**values, contact=contact).model_dump(mode='json')
        except ValidationError as exc:
            labels = []
            for error in exc.errors():
                field = '_'.join(str(x) for x in error['loc'])
                labels.append(LABELS.get(field, 'Souřadnice / hodnoty'))
            entry.update(status='error', message='Zkontrolujte pole: ' + ', '.join(dict.fromkeys(labels)))
        except (ValueError, TypeError):
            entry.update(status='error', message='Zkontrolujte počet buněk, souřadnice a datum (RRRR-MM-DD nebo DD.MM.RRRR).')
        result.append(entry)
    if not result:
        raise HTTPException(422, 'Tabulka neobsahuje žádné instituce.')
    return result, [h for h, f in zip(headers, fields) if h and not f]


async def preview_import(db, data):
    rows, ignored = await run_in_threadpool(parse_sheet, data)
    existing = (await db.execute(select(Lead.id, Lead.name, Lead.city).order_by(Lead.id))).all()
    keys = {(normalized(name), normalized(city)): lead_id for lead_id, name, city in existing}
    for row in rows:
        if row['status'] == 'error':
            continue
        key = (normalized(row['name']), normalized(row['city']))
        if key in keys:
            row.update(status='duplicate', message='Stejný název a město už existuje v CRM (i archiv) nebo dříve v souboru.', match=keys[key])
        else:
            keys[key] = f'row:{row["row"]}'
    digest = hashlib.sha256(json.dumps(rows, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    report = {'total': len(rows), 'new': sum(r['status'] == 'new' for r in rows),
              'duplicates': sum(r['status'] == 'duplicate' for r in rows),
              'errors': sum(r['status'] == 'error' for r in rows), 'ignored_columns': ignored,
              'preview_digest': digest, 'rows': [{k: v for k, v in r.items() if k not in ('data', 'match')} for r in rows]}
    return report, rows


async def apply_import(db, data, actor):
    # Serialize simultaneous imports across workers without adding tables or changing schema.
    if db.get_bind().dialect.name == 'postgresql':
        await db.execute(text('SELECT pg_advisory_xact_lock(731884021)'))
    report, rows = await preview_import(db, data)
    if not data.preview_digest or data.preview_digest != report['preview_digest']:
        raise HTTPException(409, 'Soubor nebo seznam institucí se změnil. Nejprve znovu načtěte náhled.')
    inserted = 0
    for row in rows:
        if row['status'] != 'new':
            continue
        values = LeadInput.model_validate(row['data']).model_dump(exclude={'contact'})
        lead = Lead(id=str(uuid.uuid4()), **values, updated_by=actor['actor_id'])
        db.add(lead)
        await db.flush()
        record_changes(db, actor, 'lead', lead.id, {}, values, lead.id, 'import')
        await save_contact(db, lead.id, row['data']['contact'], actor)
        inserted += 1
    return {**report, 'imported': inserted}
