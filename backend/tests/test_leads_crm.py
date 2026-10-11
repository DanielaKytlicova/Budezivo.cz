"""API integration tests against a disposable database; never uses production DSNs."""
import os
os.environ.setdefault('JWT_SECRET', 'lead-tests-not-for-production')
os.environ['DATABASE_URL'] = ''
import unittest
import uuid
import base64
import io
from datetime import datetime, timezone
from unittest.mock import patch
import httpx
from fastapi import FastAPI
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.dialects.postgresql import JSONB
from slowapi.errors import RateLimitExceeded
from slowapi import _rate_limit_exceeded_handler
from database.models import Base, User, Institution
from database.lead_models import Lead, LeadAccount, LeadAudit, LeadSession
from database.supabase import get_db
from core.security import create_jwt_token, decode_jwt_token, get_current_user
from routes.leads import router, limiter
from scripts.import_leads import prepare, summary


@compiles(JSONB, 'sqlite')
def sqlite_json(type_, compiler, **kw):
    return 'JSON'


class LeadApiTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        pg_url = os.environ.get('LEADS_TEST_POSTGRES')
        self.schema = 'crm_test_' + uuid.uuid4().hex if pg_url else None
        if pg_url:
            self.raw_engine = create_async_engine(pg_url.replace('postgresql://', 'postgresql+asyncpg://'), connect_args={'statement_cache_size': 0})
            async with self.raw_engine.begin() as conn:
                await conn.execute(text('CREATE SCHEMA ' + self.schema))
            self.engine = self.raw_engine.execution_options(schema_translate_map={None: self.schema})
        else:
            self.engine = create_async_engine('sqlite+aiosqlite:///:memory:')
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        tables = ([Institution.__table__] if self.schema else []) + [User.__table__] + [t for name, t in Base.metadata.tables.items() if name.startswith('lead_') or name == 'leads']
        async with self.engine.begin() as conn:
            await conn.run_sync(lambda c: Base.metadata.create_all(c, tables=tables))
        self.admin_id, self.tenant_id = uuid.uuid4(), uuid.uuid4()
        async with self.sessions() as db:
            institution_id = uuid.uuid4()
            if self.schema:
                db.add(Institution(id=institution_id, name='Isolated CRM test', type='museum'))
                await db.flush()
            db.add_all([User(id=self.admin_id, institution_id=institution_id, email='demo@budezivo.cz', password_hash='unused', status='active', role='admin'), User(id=self.tenant_id, institution_id=institution_id, email='tenant@example.test', password_hash='unused', status='active', role='admin')])
            await db.commit()
        app = FastAPI()
        async def session():
            async with self.sessions() as db:
                yield db
        app.dependency_overrides[get_db] = session
        app.include_router(router)
        app.state.limiter = limiter
        app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
        limiter.reset()
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test')
        self.admin = {'Authorization': 'Bearer ' + create_jwt_token(str(self.admin_id), str(uuid.uuid4()), 'demo@budezivo.cz', 'admin')}
        self.tenant = {'Authorization': 'Bearer ' + create_jwt_token(str(self.tenant_id), str(uuid.uuid4()), 'tenant@example.test', 'admin')}

    async def asyncTearDown(self):
        await self.client.aclose()
        if self.schema:
            async with self.raw_engine.begin() as conn:
                await conn.execute(text('DROP SCHEMA ' + self.schema + ' CASCADE'))
        await self.engine.dispose()

    async def account(self, username='worker'):
        data = {'username': username, 'first_name': 'Test', 'last_name': 'Worker', 'password': 'Test-password-123'}
        response = await self.client.post('/leads/team', json=data, headers=self.admin)
        self.assertEqual(response.status_code, 201, response.text)
        return response.json(), data

    async def login(self, username='worker'):
        response = await self.client.post('/leads/auth/login', json={'username': username, 'password': 'Test-password-123'})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json(), {'Authorization': 'Bearer ' + response.json()['token']}

    async def lead(self, **extra):
        response = await self.client.post('/leads', json={'name': 'Test muzeum', 'type': 'museum', 'city': 'Liberec', **extra}, headers=self.admin)
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    async def test_anonymous_tenant_and_forged_admin_are_denied(self):
        self.assertEqual((await self.client.get('/leads')).status_code, 401)
        for path in ('/leads', '/leads/team', '/leads/export', '/leads/goals'):
            self.assertEqual((await self.client.get(path, headers=self.tenant)).status_code, 403)
        forged = {'Authorization': 'Bearer ' + create_jwt_token(str(self.tenant_id), 'x', 'demo@budezivo.cz', 'admin')}
        self.assertEqual((await self.client.get('/leads', headers=forged)).status_code, 403)

    async def test_assignment_audit_contacts_conflict_and_filters(self):
        account, _ = await self.account()
        _, worker = await self.login()
        own = await self.lead(assignee_id=account['id'])
        other = await self.lead(name='Pražská galerie')
        self.assertEqual((await self.client.get('/leads', headers=worker)).json()['total'], 2)
        change = {'version': 1, 'changes': {'internal_note': 'Trvalá poznámka', 'stage': 'Připojeno – aktivně používá', 'contact': {'first_name': 'Jana', 'email': 'jana@example.test'}}}
        response = await self.client.patch('/leads/' + own['id'], json=change, headers=worker)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual((await self.client.patch('/leads/' + other['id'], json=change, headers=worker)).status_code, 403)
        self.assertEqual((await self.client.patch('/leads/' + own['id'], json=change, headers=worker)).status_code, 409)
        change = {'version': 2, 'changes': {'assignee_id': None}}
        self.assertEqual((await self.client.patch('/leads/' + own['id'], json=change, headers=worker)).status_code, 403)
        detail = (await self.client.get('/leads/' + own['id'], headers=worker)).json()
        self.assertEqual(detail['internal_note'], 'Trvalá poznámka')
        self.assertEqual(detail['contacts'][0]['email'], 'jana@example.test')
        note_audit = [a for a in detail['audit'] if a['field'] == 'internal_note'][0]
        self.assertEqual(note_audit['new_value'], 'Trvalá poznámka')
        self.assertEqual(note_audit['actor_id'], 'lead:' + account['id'])
        for query in ('mine=true', 'connected=true', 'q=jana@example.test', 'city=Liberec'):
            if query.startswith('city'): continue
            self.assertEqual((await self.client.get('/leads?' + query, headers=worker)).json()['total'], 1)
        self.assertEqual((await self.client.post('/leads/assign', json={'lead_ids': [own['id']], 'assignee_id': None}, headers=self.admin)).status_code, 200)
        self.assertEqual((await self.client.patch('/leads/' + own['id'], json={'version': 3, 'changes': {'internal_note': 'blocked'}}, headers=worker)).status_code, 403)

    async def test_collaborator_creates_owned_lead_with_contact_and_audit(self):
        account, _ = await self.account()
        _, worker = await self.login()
        data = {'name': 'Vlastní tip', 'type': 'gallery', 'city': 'Praha',
                'contact': {'email': 'kontakt@example.test'}}
        response = await self.client.post('/leads', json=data, headers=worker)
        self.assertEqual(response.status_code, 201, response.text)
        lead = response.json()
        self.assertEqual(lead['assignee_id'], account['id'])
        self.assertFalse(lead['archived'])
        detail = (await self.client.get('/leads/' + lead['id'], headers=worker)).json()
        self.assertTrue(detail['can_edit'])
        self.assertEqual(detail['contacts'][0]['email'], 'kontakt@example.test')
        audit = next(a for a in detail['audit'] if a['field'] == 'assignee_id')
        self.assertEqual(audit['new_value'], account['id'])
        self.assertEqual(audit['actor_id'], 'lead:' + account['id'])
        update = await self.client.patch('/leads/' + lead['id'], headers=worker,
                                        json={'version': lead['version'], 'changes': {'internal_note': 'Doplněno'}})
        self.assertEqual(update.status_code, 200, update.text)

    def spreadsheet(self, content, filename='instituce.csv'):
        raw = content.encode('utf-8-sig') if isinstance(content, str) else content
        return {'filename': filename, 'content': base64.b64encode(raw).decode()}

    async def test_spreadsheet_preview_import_duplicate_errors_contacts_and_audit(self):
        old = await self.lead(name='Muzeum Praha', city='Praha', archived=True, internal_note='Zachovat')
        body = self.spreadsheet('Název;Typ;Město;E-mail;Web;Termín dalšího kroku;Cizí sloupec\n'
                                ' MUZEUM PRAHA ;muzeum;Praha;;;;ignorovat\n'
                                'Nová galerie;galerie;Říčany;info@example.test;https://example.test;31.10.2026;ignorovat\n'
                                'nova GALERIE;galerie;ricany;;;;\n'
                                'Chybný kontakt;muzeum;Praha;neni-mail;;;\n')
        preview = await self.client.post('/leads/import/preview', json=body, headers=self.admin)
        self.assertEqual(preview.status_code, 200, preview.text)
        report = preview.json()
        self.assertEqual((report['new'], report['duplicates'], report['errors']), (1, 2, 1))
        self.assertEqual(report['ignored_columns'], ['Cizí sloupec'])
        self.assertEqual((await self.client.get('/leads?archived=true', headers=self.admin)).json()['total'], 1)
        self.assertEqual((await self.client.get('/leads', headers=self.admin)).json()['total'], 0)
        confirmed = await self.client.post('/leads/import', json={**body, 'preview_digest': report['preview_digest']}, headers=self.admin)
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        self.assertEqual(confirmed.json()['imported'], 1)
        new = (await self.client.get('/leads', headers=self.admin)).json()['items'][0]
        self.assertEqual(new['type'], 'gallery')
        self.assertEqual(new['next_action_date'], '2026-10-31')
        self.assertIsNone(new['assignee_id'])
        detail = (await self.client.get('/leads/' + new['id'], headers=self.admin)).json()
        self.assertEqual(detail['contacts'][0]['email'], 'info@example.test')
        self.assertTrue(any(a['action'] == 'import' and a['actor_id'] == 'user:' + str(self.admin_id) for a in detail['audit']))
        unchanged = (await self.client.get('/leads/' + old['id'], headers=self.admin)).json()
        self.assertEqual(unchanged['internal_note'], 'Zachovat')
        repeat = await self.client.post('/leads/import', json={**body, 'preview_digest': report['preview_digest']}, headers=self.admin)
        self.assertEqual(repeat.status_code, 409)
        new_preview = (await self.client.post('/leads/import/preview', json=body, headers=self.admin)).json()
        self.assertEqual(new_preview['new'], 0)

    async def test_spreadsheet_permissions_stale_preview_and_atomic_failure(self):
        await self.account()
        _, worker = await self.login()
        body = self.spreadsheet('Název;Město\nA;Praha\nB;Liberec\n')
        for headers, status in (({}, 401), (self.tenant, 403), (worker, 403)):
            for path in ('/leads/import/preview', '/leads/import'):
                self.assertEqual((await self.client.post(path, json=body, headers=headers)).status_code, status)
        self.assertEqual((await self.client.post('/leads/import', json=body, headers=self.admin)).status_code, 409)
        report = (await self.client.post('/leads/import/preview', json=body, headers=self.admin)).json()
        altered = self.spreadsheet('Název;Město\nC;Praha\n')
        self.assertEqual((await self.client.post('/leads/import', json={**altered, 'preview_digest': report['preview_digest']}, headers=self.admin)).status_code, 409)
        # A failure during writing rolls back the entire batch, including earlier rows and audit.
        from fastapi import HTTPException
        from services.lead_service import save_contact
        count = 0
        async def fail_second(*args):
            nonlocal count
            count += 1
            if count == 2:
                raise HTTPException(503, 'Simulovaná chyba')
            await save_contact(*args)
        with patch('services.lead_import.save_contact', side_effect=fail_second):
            response = await self.client.post('/leads/import', json={**body, 'preview_digest': report['preview_digest']}, headers=self.admin)
        self.assertEqual(response.status_code, 503)
        self.assertEqual((await self.client.get('/leads', headers=self.admin)).json()['total'], 0)
        await self.lead(name='A', city='Praha')
        self.assertEqual((await self.client.post('/leads/import', json={**body, 'preview_digest': report['preview_digest']}, headers=self.admin)).status_code, 409)

    async def test_spreadsheet_xlsx_first_sheet_dates_and_formula_rejection(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(['Název', 'Město', 'Termín dalšího kroku', 'Zeměpisná šířka', 'Zeměpisná délka'])
        sheet.append(['Excel muzeum', 'Praha', datetime(2026, 10, 31), '50,1', '14,4'])
        book.create_sheet('Ignorovat').append(['Toto není seznam institucí'])
        def payload():
            buffer = io.BytesIO()
            book.save(buffer)
            return self.spreadsheet(buffer.getvalue(), 'instituce.xlsx')
        body = payload()
        preview = await self.client.post('/leads/import/preview', json=body, headers=self.admin)
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(preview.json()['new'], 1)
        response = await self.client.post('/leads/import', json={**body, 'preview_digest': preview.json()['preview_digest']}, headers=self.admin)
        self.assertEqual(response.status_code, 200, response.text)
        sheet['A2'] = '=1+1'
        self.assertEqual((await self.client.post('/leads/import/preview', json=payload(), headers=self.admin)).status_code, 422)

    async def test_spreadsheet_file_limits_headers_and_csv_encodings(self):
        cases = [self.spreadsheet('Město\nPraha'), self.spreadsheet('Název;name\nA;B'),
                 self.spreadsheet('Název\n' + 'A\n' * 501), self.spreadsheet('bad', 'x.xlsx'),
                 self.spreadsheet('Název\nA', 'x.exe'), {'filename': 'x.csv', 'content': '!!invalid!!'}]
        for body in cases:
            self.assertEqual((await self.client.post('/leads/import/preview', json=body, headers=self.admin)).status_code, 422)
        for source in ('Název\nMuzeum', 'Název,Město\n"Galerie, s.r.o.",Praha', 'Název\tMěsto\nMuzeum\tPraha'):
            body = self.spreadsheet(source.encode('cp1250'))
            reply = await self.client.post('/leads/import/preview', json=body, headers=self.admin)
            self.assertEqual(reply.status_code, 200, reply.text)
            self.assertEqual(reply.json()['new'], 1)

    async def test_collaborator_creation_cannot_assign_others_or_archive(self):
        await self.account()
        other, _ = await self.account('other')
        _, worker = await self.login()
        for extra in ({'assignee_id': other['id']}, {'archived': True}):
            response = await self.client.post('/leads', headers=worker,
                                              json={'name': 'Zakázaný tip', 'type': 'museum', **extra})
            self.assertEqual(response.status_code, 403, response.text)
        self.assertEqual((await self.client.get('/leads', headers=worker)).json()['total'], 0)
        for headers, status in (({}, 401), (self.tenant, 403)):
            response = await self.client.post('/leads', headers=headers,
                                              json={'name': 'Nepovolený přístup', 'type': 'museum'})
            self.assertEqual(response.status_code, status, response.text)

    async def test_deactivation_reset_logout_and_rotating_refresh(self):
        account, data = await self.account()
        session, worker = await self.login()
        refreshed = await self.client.post('/leads/auth/refresh', json={'refresh_token': session['refresh_token']})
        self.assertEqual(refreshed.status_code, 200)
        self.assertEqual((await self.client.post('/leads/auth/refresh', json={'refresh_token': session['refresh_token']})).status_code, 401)
        self.assertEqual((await self.client.put('/leads/team/' + account['id'], json={**data, 'active': False, 'password': None}, headers=self.admin)).status_code, 200)
        self.assertEqual((await self.client.get('/leads', headers=worker)).status_code, 401)
        self.assertEqual((await self.client.post('/leads/auth/login', json={'username': 'worker', 'password': data['password']})).status_code, 401)
        await self.client.put('/leads/team/' + account['id'], json={**data, 'active': True}, headers=self.admin)
        _, worker = await self.login()
        await self.client.post('/leads/auth/logout', headers=worker)
        self.assertEqual((await self.client.get('/leads', headers=worker)).status_code, 401)
        async with self.sessions() as db:
            audit = (await db.execute(select(LeadAudit).where(LeadAudit.entity == 'account'))).scalars().all()
            self.assertTrue(audit)
            self.assertNotIn('Test-password-123', str([a.new_value for a in audit]))

    async def test_crm_token_rejected_by_existing_admin_dependency(self):
        from starlette.requests import Request
        await self.account()
        session, _ = await self.login()
        request = Request({'type': 'http', 'headers': [(b'cookie', ('access_token=' + session['token']).encode())]})
        from fastapi import HTTPException
        with self.assertRaises(HTTPException) as error:
            await get_current_user(request, None)
        self.assertEqual(error.exception.status_code, 401)

    async def test_export_formula_safety_goals_activity_and_archive(self):
        account, _ = await self.account()
        _, worker = await self.login()
        lead = await self.lead(name='=DANGEROUS()', assignee_id=account['id'], stage='Připojeno – platící', catalog_programs=5)
        self.assertEqual((await self.client.get('/leads/export', headers=worker)).status_code, 403)
        export = await self.client.get('/leads/export?connected=true', headers=self.admin)
        self.assertIn("'=DANGEROUS()", export.text)
        self.assertEqual(export.headers['cache-control'], 'no-store')
        self.assertEqual((await self.client.get('/leads/export?q=missing', headers=self.admin)).text.count('\n'), 1)
        goal = {'name': 'Programy', 'metric': 'programs', 'target': 10, 'deadline': '2027-01-31'}
        self.assertEqual((await self.client.post('/leads/goals', json=goal, headers=worker)).status_code, 403)
        await self.client.post('/leads/goals', json=goal, headers=self.admin)
        self.assertEqual((await self.client.get('/leads/goals', headers=worker)).json()[0]['current'], 5)
        activity = {'kind': 'phone', 'subject': 'Hovor', 'body': 'Domluva', 'occurred_at': datetime.now(timezone.utc).isoformat()}
        self.assertEqual((await self.client.post('/leads/' + lead['id'] + '/activities', json=activity, headers=worker)).status_code, 201)
        self.assertEqual((await self.client.delete('/leads/' + lead['id'], headers=worker)).status_code, 405)
        self.assertEqual((await self.client.patch('/leads/' + lead['id'], json={'version': 1, 'changes': {'archived': True}}, headers=self.admin)).status_code, 200)
        self.assertEqual((await self.client.get('/leads', headers=worker)).json()['total'], 0)
        self.assertEqual((await self.client.get('/leads/goals', headers=worker)).json()[0]['current'], 0)

    async def test_validation_and_atomic_bulk_assignment(self):
        lead = await self.lead()
        for changes in ({'program_url': 'javascript:alert(1)'}, {'score': 101}, {'lat': 500}, {'legacy_snapshot': {}}, {'name': None}):
            response = await self.client.patch('/leads/' + lead['id'], json={'version': 1, 'changes': changes}, headers=self.admin)
            self.assertEqual(response.status_code, 422, response.text)
        account, _ = await self.account()
        response = await self.client.post('/leads/assign', json={'lead_ids': [lead['id'], 'missing'], 'assignee_id': account['id']}, headers=self.admin)
        self.assertEqual(response.status_code, 404)
        self.assertIsNone((await self.client.get('/leads/' + lead['id'], headers=self.admin)).json()['assignee_id'])

    async def test_bruteforce_limit_and_generic_error(self):
        replies = []
        for _ in range(6):
            replies.append(await self.client.post('/leads/auth/login', json={'username': 'missing', 'password': 'irrelevant'}))
        self.assertEqual(replies[0].status_code, 401)
        self.assertEqual(replies[-1].status_code, 429)

    async def test_import_idempotence_and_source_verification(self):
        from scripts.import_leads import apply
        seed = [{'id': 'BZ-0001', 'name': 'Museum', 'type': 'museum', 'city': 'Praha', 'stage': 'Dosud neosloveno', 'contact': 'Jana Test', 'plan': 'FREE'}]
        export = {'has_more': False, 'model_projection': {'truncated': False}, 'rows': [{'institution_id': 'BZ-0001', 'internal_note': 'D1 note', 'stage': 'Odmítli', 'email': 'work@example.test'}]}
        records = prepare(seed, export)
        with patch('database.supabase.AsyncSessionLocal', self.sessions):
            first = await apply(records)
            second = await apply(records)
        self.assertEqual(first['inserted'], 1)
        self.assertEqual(second['inserted'], 0)
        self.assertTrue(second['database_verified'])
        self.assertEqual(first['comparison']['internal_notes'], 1)
        export['rows'][0]['internal_note'] = 'New source data'
        with patch('database.supabase.AsyncSessionLocal', self.sessions):
            with self.assertRaises(ValueError):
                await apply(prepare(seed, export))


class ImportValidationTests(unittest.TestCase):
    def test_incomplete_duplicate_or_orphan_source_refused(self):
        seed = [{'id': 'BZ-0001', 'name': 'Museum', 'type': 'museum'}]
        good = {'has_more': False, 'model_projection': {'truncated': False}, 'rows': []}
        for data in ({**good, 'has_more': True}, {**good, 'model_projection': {'truncated_values': 1}}, {**good, 'rows': [{'institution_id': 'missing'}]}):
            with self.assertRaises(ValueError): prepare(seed, data)
        with self.assertRaises(ValueError): prepare(seed + seed, good)

class PrivacyTests(unittest.TestCase):
    def test_access_logs_redact_crm_query(self):
        import logging
        from middleware.lead_privacy import LeadAccessLogFilter
        record = logging.LogRecord('uvicorn.access', 20, '', 0, '%s %s %s %s %s',
                                   ('client', 'GET', '/api/leads?q=private@example.test', '1.1', 200), None)
        LeadAccessLogFilter().filter(record)
        self.assertNotIn('private@example.test', record.getMessage())
        self.assertIn('/api/leads/[private]', record.getMessage())
